"""T13 Excel 导入导出底座（对齐 Go modules/dataio 全部语义）。

模块注册表模式：新模块实现 ModuleSpec 即可挂入（funding.income 首个全链路落地）。
- 模板：表头 + 示例行 + 「校验说明」sheet（openpyxl 对齐 excelize 行为）
- 导入：xlsx 逐行解析 → 幂等键去重跳过 → 行级错误清单 → sys_import_log snapshot
- 撤销：snapshot 内业务行未被引用才允许（引用中 → 40902 拒绝整体撤销）
- 导出：选择性字段（columns/fields）与顺序，缺省全字段
"""
from __future__ import annotations

import io
import json
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Any, Callable

from django.http import HttpRequest
from django.utils import timezone as dj_tz
from openpyxl import Workbook, load_workbook

from apps.asset.models import AAsset, AStockTxn
from apps.funding.models import FIncome
from apps.notify.models import SysImportLog
from core.auth.deps import domain_from
from core.httpx import ApiError, CODE_BAD_REQUEST, CODE_CONFLICT, CODE_NOT_FOUND, ok

MOD_INCOME = "funding.income"
MOD_INVENTORY = "asset.inventory"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@dataclass
class ModuleSpec:
    module: str
    headers: list[str]
    sample: list[str]
    hints: list[str]
    export_fields: list[str]
    export_labels: dict[str, str]
    # rows → (imported, skipped, errors, snapshot_keys[, summary dict])
    import_fn: Callable[..., tuple]
    # snapshot_keys → 被业务引用的行数（>0 禁止撤销）
    claimed_fn: Callable[[int, list[str]], int] = field(default=lambda t, k: 0)
    # snapshot_keys → 撤销（返回删除行数）
    undo_fn: Callable[[int, list[str]], int] = field(default=lambda t, k: 0)
    # (kw, claim_status) → 导出行记录 dict（key=export_fields）
    export_fn: Callable[[int, str | None, str | None], list[dict]] = field(
        default=lambda t, kw, cs: []
    )


def _parse_amount(s: str):
    s = (s or "").strip()
    if not s:
        return None, False
    try:
        return Decimal(s).quantize(Decimal("0.01")), True
    except InvalidOperation:
        return None, False


def _parse_date(s: str):
    from datetime import date, datetime

    s = (s or "").strip()
    if not s:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    if isinstance(s, date):
        return s
    return None


# ---------------------------------------------------------------------------
# funding.income：收款导入（幂等：bank_flow_no 同租户唯一，重复行跳过）
# ---------------------------------------------------------------------------
def _income_import(tenant_id: int, user_id: int, rows: list[list[str]]):
    imported = skipped = 0
    errs: list[dict] = []
    snapshot: list[str] = []
    for i, row in enumerate(rows):
        line_no = i + 2  # xlsx 行号（含表头）
        if len(row) < 6:
            errs.append({"row": line_no, "error": "列数不足 6"})
            continue
        flow_no = (row[5] or "").strip()
        if not flow_no:
            errs.append({"row": line_no, "error": "银行流水号必填"})
            continue
        amt, ok_amt = _parse_amount(row[2])
        if not ok_amt:
            errs.append({"row": line_no, "error": f"金额非法：{row[2]}"})
            continue
        if FIncome.objects.filter(tenant_id=tenant_id, bank_flow_no=flow_no, deleted=0).exists():
            skipped += 1
            continue
        FIncome.objects.create(
            tenant_id=tenant_id, creator_id=user_id, legal_entity_id=1,
            receipt_no=(row[0] or "").strip() or None,
            receipt_date=_parse_date(row[1]),
            amount=amt,
            payer_name=(row[3] or "").strip() or None,
            purpose_text=(row[4] or "").strip() or None,
            bank_flow_no=flow_no,
        )
        imported += 1
        snapshot.append(flow_no)
    return imported, skipped, errs, snapshot


def _income_claimed(tenant_id: int, flow_nos: list[str]) -> int:
    return FIncome.objects.filter(
        tenant_id=tenant_id, bank_flow_no__in=flow_nos, claim_status="CLAIMED", deleted=0
    ).count()


def _income_undo(tenant_id: int, flow_nos: list[str]) -> int:
    return FIncome.objects.filter(
        tenant_id=tenant_id, bank_flow_no__in=flow_nos, deleted=0
    ).update(deleted=1, updated_at=dj_tz.now())


INCOME_EXPORT_FIELDS = [
    "id", "code", "receipt_no", "receipt_date", "amount",
    "payer_name", "purpose_text", "bank_flow_no", "claim_status", "created_at",
]
INCOME_EXPORT_LABEL = {
    "id": "ID", "code": "到款编号", "receipt_no": "收款编号", "receipt_date": "收款日期",
    "amount": "金额", "payer_name": "付款单位", "purpose_text": "摘要",
    "bank_flow_no": "银行流水号", "claim_status": "认领状态", "created_at": "创建时间",
}


def _income_export(tenant_id: int, kw: str | None, claim_status: str | None) -> list[dict]:
    qs = FIncome.objects.filter(tenant_id=tenant_id, deleted=0)
    if kw:
        from django.db.models import Q
        qs = qs.filter(Q(payer_name__icontains=kw) | Q(purpose_text__icontains=kw) | Q(bank_flow_no__icontains=kw))
    if claim_status:
        qs = qs.filter(claim_status=claim_status)
    out = []
    for e in qs.order_by("-id")[:10000]:
        out.append({
            "id": str(e.id), "code": e.code or "",
            "receipt_no": e.receipt_no or "",
            "receipt_date": f"{e.receipt_date:%Y-%m-%d}" if e.receipt_date else "",
            "amount": f"{e.amount:.2f}", "payer_name": e.payer_name or "",
            "purpose_text": e.purpose_text or "", "bank_flow_no": e.bank_flow_no or "",
            "claim_status": e.claim_status,
            "created_at": f"{dj_tz.localtime(e.created_at):%Y-%m-%d %H:%M:%S}",
        })
    return out


# ---------------------------------------------------------------------------
# asset.inventory：盘点导入（A-11 三方比对语义：实物采集行 vs 资产卡）
#   · 序列号命中台账（在册非处置）→ 匹配成功
#   · 实物有、账上无 → 盘盈 COUNT_GAIN 流水（snapshot 可撤销）
#   · 账上有、实物未采集 → 盘亏 COUNT_LOSS 流水（snapshot 可撤销）
# ---------------------------------------------------------------------------
INVENTORY_HEADERS = ["资产编号", "序列号", "资产名称", "存放地点", "数量", "盘点人", "盘点日期", "备注"]
INVENTORY_FIELDS = ["code", "serial_no", "name", "location", "qty", "custodian", "status", "value"]


def _parse_qty(s: str):
    s = (s or "").strip()
    if not s:
        return Decimal(1), True  # 缺省 1 件
    try:
        q = Decimal(s)
        return (q, True) if q > 0 else (None, False)
    except InvalidOperation:
        return None, False


def _inventory_txn(tenant_id: int, user_id: int, type_: str, **kw) -> str:
    import uuid
    t = AStockTxn.objects.create(
        tenant_id=tenant_id, creator_id=user_id, operator_id=user_id,
        legal_entity_id=1, security_level="INTERNAL", type=type_, status="ISSUED",
        code=f"TXN-INV-{dj_tz.now():%Y%m%d%H%M%S}-{uuid.uuid4().hex[:6]}",
        **kw,
    )
    return t.code


def _inventory_import(tenant_id: int, user_id: int, rows: list[list[str]]):
    from django.db.models import Q

    imported = matched = gains = losses = skipped = 0
    errs: list[dict] = []
    snapshot: list[str] = []
    seen: set[str] = set()
    # 台账侧：在册（非处置/待处置）且有序列号的卡
    active = {
        a.serial_no: a
        for a in AAsset.objects.filter(tenant_id=tenant_id, deleted=0)
        .exclude(status__in=("DISPOSED", "PENDING_DISPOSE"))
        .exclude(serial_no__isnull=True).exclude(serial_no="")
    }
    collected: set[str] = set()
    for i, row in enumerate(rows):
        line_no = i + 2  # xlsx 行号（含表头）
        if len(row) < 8:
            row = row + [""] * (8 - len(row))
        serial = (row[1] or "").strip()
        if not serial:
            errs.append({"row": line_no, "error": "序列号必填（三方比对匹配键）"})
            continue
        qty, ok_qty = _parse_qty(row[4])
        if not ok_qty:
            errs.append({"row": line_no, "error": f"数量非法：{row[4]}"})
            continue
        if serial in seen:
            skipped += 1  # 文件内重复行（幂等）
            continue
        seen.add(serial)
        collected.add(serial)
        a = active.get(serial)
        if a is not None:
            matched += 1  # 账实相符（位置差异仅提示，不落流水，对齐 A-11）
            continue
        remark = "盘点导入盘盈：" + "，".join(
            x for x in [f"编号{row[0]}" if row[0].strip() else "",
                        f"名称{row[2]}" if row[2].strip() else "",
                        f"地点{row[3]}" if row[3].strip() else "",
                        f"盘点人{row[5]}" if row[5].strip() else "",
                        f"日期{row[6]}" if row[6].strip() else ""] if x
        ) or "实物在册但账上无卡"
        snapshot.append(_inventory_txn(
            tenant_id, user_id, "COUNT_GAIN",
            ref_sku=serial[:64], qty=qty, balance_after=qty, remark=remark[:500],
        ))
        gains += 1
    # 账实核对：账上在册卡未被实物采集到 → 盘亏
    for a in active.values():
        if a.serial_no not in collected:
            snapshot.append(_inventory_txn(
                tenant_id, user_id, "COUNT_LOSS",
                ref_asset_id=a.id, qty=1, balance_after=0,
                remark=f"盘点导入盘亏：{a.code or a.id} {a.name}",
            ))
            losses += 1
    imported = matched + gains
    summary = {"matchedCount": matched, "gainCount": gains, "lossCount": losses}
    return imported, skipped, errs, snapshot, summary


def _inventory_claimed(tenant_id: int, codes: list[str]) -> int:
    return AStockTxn.objects.filter(
        tenant_id=tenant_id, code__in=codes, deleted=0
    ).exclude(status="ISSUED").count()


def _inventory_undo(tenant_id: int, codes: list[str]) -> int:
    return AStockTxn.objects.filter(
        tenant_id=tenant_id, code__in=codes, deleted=0
    ).update(deleted=1, updated_at=dj_tz.now())


def _inventory_export(tenant_id: int, kw: str | None, claim_status: str | None) -> list[dict]:
    qs = AAsset.objects.filter(tenant_id=tenant_id, deleted=0)
    if kw:
        from django.db.models import Q
        qs = qs.filter(Q(name__icontains=kw) | Q(serial_no__icontains=kw) | Q(code__icontains=kw))
    out = []
    for a in qs.order_by("-id")[:10000]:
        out.append({
            "code": a.code or "", "serial_no": a.serial_no or "", "name": a.name or "",
            "location": a.location or "", "qty": "1", "custodian": "",
            "status": a.status, "value": f"{a.value:.2f}",
        })
    return out


SPECS: dict[str, ModuleSpec] = {
    MOD_INCOME: ModuleSpec(
        module=MOD_INCOME,
        headers=["收款编号", "收款日期", "金额", "付款单位", "摘要", "银行流水号"],
        sample=["", "2026-01-15", "50000.00", "示例科技有限公司", "纵向课题 2026 年度拨款", "BK2026011500001"],
        hints=[
            "收款日期格式 YYYY-MM-DD；金额两位小数",
            "银行流水号必填且同租户唯一（重复行自动跳过，幂等）",
            "导入成功后可在「导入日志」中整体撤销（未认领状态下）",
        ],
        export_fields=INCOME_EXPORT_FIELDS,
        export_labels=INCOME_EXPORT_LABEL,
        import_fn=_income_import,
        claimed_fn=_income_claimed,
        undo_fn=_income_undo,
        export_fn=_income_export,
    ),
    MOD_INVENTORY: ModuleSpec(
        module=MOD_INVENTORY,
        headers=INVENTORY_HEADERS,
        sample=["AST-DEMO-0001", "SN-DEMO-0001", "X 射线衍射仪（示例）", "分析所 101 室", "1", "示例盘点人", "2026-10-15", "在库可用"],
        hints=[
            "序列号必填：与资产卡序列号匹配（三方比对匹配键）",
            "数量缺省 1；文件内重复序列号自动跳过（幂等）",
            "实物有而账上无 → 盘盈流水；账上有而未采集 → 盘亏流水（均可在导入日志整体撤销）",
            "已处置/待处置资产不参与比对",
        ],
        export_fields=INVENTORY_FIELDS,
        export_labels={"code": "资产编号", "serial_no": "序列号", "name": "资产名称",
                        "location": "存放地点", "qty": "数量", "custodian": "保管人",
                        "status": "状态", "value": "原值"},
        import_fn=_inventory_import,
        claimed_fn=_inventory_claimed,
        undo_fn=_inventory_undo,
        export_fn=_inventory_export,
    ),
}


def spec_or_400(module: str) -> ModuleSpec:
    spec = SPECS.get(module or "")
    if spec is None:
        raise ApiError(400, CODE_BAD_REQUEST, "未知模块，支持：" + " | ".join(SPECS))
    return spec


# ---------------------------------------------------------------------------
# xlsx 工具（openpyxl）
# ---------------------------------------------------------------------------
def build_template(spec: ModuleSpec) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "导入模板"
    ws.append(spec.headers)
    ws.append(spec.sample)
    hints = wb.create_sheet("校验说明")
    for h in spec.hints:
        hints.append([h])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_body_rows(file_obj) -> list[list[str]]:
    """首 sheet 除表头外的非空行（单元格统一转字符串）。"""
    try:
        wb = load_workbook(file_obj, read_only=True, data_only=True)
    except Exception:
        raise ApiError(400, CODE_BAD_REQUEST, "仅支持 xlsx 格式") from None
    ws = wb.worksheets[0]
    body: list[list[str]] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:
            continue
        cells = ["" if v is None else (v if isinstance(v, str) else str(v)) for v in row]
        if any(c.strip() for c in cells):
            body.append(cells)
    wb.close()
    if not body:
        raise ApiError(400, CODE_BAD_REQUEST, "无数据行")
    return body


def build_export(spec: ModuleSpec, rows: list[dict], columns: list[str]) -> bytes:
    valid = set(spec.export_fields)
    cols = [c for c in columns if c in valid] or list(spec.export_fields)
    wb = Workbook()
    ws = wb.active
    ws.title = "导出结果"
    ws.append([spec.export_labels.get(c, c) for c in cols])
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
